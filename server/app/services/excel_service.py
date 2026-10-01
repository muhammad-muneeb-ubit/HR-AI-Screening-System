from io import BytesIO
import re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


def create_analysis_excel(rows, job_title):
    print("Creating Excel file for job title:", job_title)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = f"{job_title} - Resume Analysis"
    headers = [
        "File Name",
        "Candidate Name",
        "Email",
        "Phone",
        "Status",
        "Score",
        "Skills Analysis",
        "Experience Analysis",
        "Qualifications Analysis",
        "Overall Response",
        "Analyzed At",
    ]

    worksheet.append(headers)

    # Header styling
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # Add data
    for row in rows:
        # print("wof ", row["score"]*100)
        # print("wf", float(row["score"]*100))
        # print("wi", int(row["score"]*100))
        worksheet.append([
            row["file_name"],
            row["candidate_name"],
            row["candidate_email"],
            row["candidate_phone"],
            row["status"],
            int(row["score"]*100) if row["score"] is not None else None,
            row["skills_analysis"],
            row["experience_analysis"],
            row["qualifications_analysis"],
            row["overall_response"],
            row["created_at"],
        ])

    # Wrap long text
    for row in worksheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    # Set reasonable column widths
    widths = {
        1: 30,   # File Name
        2: 25,   # Candidate Name
        3: 35,   # Email
        4: 25,   # Phone
        5: 12,   # Status
        6: 12,   # Score
        7: 60,   # Skills
        8: 60,   # Experience
        9: 60,   # Qualifications
        10: 80,  # Overall
        11: 25,  # Date
    }

    for column, width in widths.items():
        worksheet.column_dimensions[
            get_column_letter(column)
        ].width = width

    # Freeze header row
    worksheet.freeze_panes = "A2"

    # Auto filter
    worksheet.auto_filter.ref = worksheet.dimensions

    # Save workbook to memory
    output = BytesIO()
    workbook.save(output)

    output.seek(0)

    return output

def clean_sheet_title(title):
    title = str(title).strip()

    # Remove Excel-invalid characters
    title = re.sub(r'[\\/*?:\[\]]', '-', title)

    # Excel sheet name max length = 31
    title = title[:31]

    # Excel doesn't like an empty sheet name
    return title or "Sheet1"

